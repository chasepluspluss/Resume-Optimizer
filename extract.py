"""Turn uploaded files (or pasted text) into plain text, then sort that text into resume sections."""

import io
import os
import re
import zipfile
from html import unescape
from xml.etree import ElementTree as ET


# --------------------------------------------------------------------------- file -> text

def extract_text(filename, data):
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if ext == "pdf":
        return _from_pdf(data)
    if ext in ("docx", "docm", "dotx"):
        return _from_docx(data)
    if ext == "odt":
        return _from_odt(data)
    if ext == "rtf":
        return _from_rtf(_decode(data))
    if ext in ("html", "htm"):
        return _from_html(_decode(data))
    if ext == "doc":
        raise ValueError("old .doc format isn't supported - save it as .docx or PDF and upload again")
    if ext in ("png", "jpg", "jpeg", "gif", "webp", "heic", "bmp", "tif", "tiff"):
        raise ValueError("images can't be read (no OCR) - upload the PDF/DOCX or paste the text")
    text = _decode(data)
    if "\x00" in text:
        raise ValueError(f"unsupported file type (.{ext or '?'})")
    return text


def _decode(data):
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("latin-1")


def _from_pdf(data):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("PDF is password-protected")
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _from_docx(data):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(data))
    lines = []

    def add_paragraph(p):
        text = p.text.strip()
        if not text:
            lines.append("")
            return
        try:
            style = (p.style.name or "").lower()
        except Exception:
            style = ""
        ppr = p._p.pPr
        is_list = (ppr is not None and ppr.numPr is not None) or "list" in style
        lines.append(("• " + text) if is_list else text)

    def add_block(parent_el, parent):
        for el in parent_el.iterchildren():
            tag = el.tag.rsplit("}", 1)[-1]
            if tag == "p":
                add_paragraph(Paragraph(el, parent))
            elif tag == "tbl":
                for row in Table(el, parent).rows:
                    seen = []
                    for cell in row.cells:
                        if any(cell._tc is s for s in seen):  # merged cells repeat
                            continue
                        seen.append(cell._tc)
                        add_block(cell._tc, cell)
            elif tag == "sdt":  # content controls used by many templates
                content = el.find(qn("w:sdtContent"))
                if content is not None:
                    add_block(content, parent)

    for section in doc.sections[:1]:
        for p in section.header.paragraphs:
            add_paragraph(p)
    add_block(doc.element.body, doc)

    # Designer templates often keep everything in text boxes, which the walk above skips.
    if sum(len(l) for l in lines) < 200:
        fallback = []
        for p in doc.element.body.iter(qn("w:p")):
            if any(a.tag.endswith("}Fallback") for a in p.iterancestors()):
                continue
            fallback.append("".join(t.text or "" for t in p.iter(qn("w:t"))).strip())
        if sum(len(l) for l in fallback) > sum(len(l) for l in lines):
            lines = fallback
    return "\n".join(lines)


def _from_odt(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        root = ET.fromstring(z.read("content.xml"))
    text_ns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
    lines = []

    def walk(el, in_list):
        if el.tag in (text_ns + "p", text_ns + "h"):
            t = "".join(el.itertext()).strip()
            lines.append(("• " + t) if (in_list and t) else t)
            return
        for child in el:
            walk(child, in_list or el.tag == text_ns + "list-item")

    walk(root, False)
    return "\n".join(lines)


def _from_rtf(s):
    s = re.sub(r"\{\\(?:fonttbl|colortbl|stylesheet|info|\*)[^{}]*(?:\{[^{}]*\}[^{}]*)*\}", "", s)
    s = re.sub(r"\\'([0-9a-fA-F]{2})", lambda m: bytes([int(m.group(1), 16)]).decode("cp1252", "replace"), s)
    s = re.sub(r"\\u(-?\d+)\??", lambda m: chr(int(m.group(1)) % 65536), s)
    s = re.sub(r"\\(?:par|line)d?\b ?", "\n", s)
    s = re.sub(r"\\tab\b ?", "\t", s)
    s = re.sub(r"\\bullet\b ?", "• ", s)
    s = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", s)
    return s.replace("{", "").replace("}", "")


def _from_html(s):
    s = re.sub(r"(?is)<(script|style|head)[^>]*>.*?</\1>", "", s)
    s = re.sub(r"(?i)<li[^>]*>", "\n• ", s)
    s = re.sub(r"(?i)<br\s*/?>|</(p|div|h\d|li|tr|section|header)>", "\n", s)
    s = re.sub(r"(?i)</t[dh]>", "\t", s)
    s = re.sub(r"<[^>]+>", "", s)
    return unescape(s)


# --------------------------------------------------------------------------- text -> sections

BULLET_RE = re.compile(r"^\s*(?:[•●○◦▪▫■□‣⁃∙·➢➤►▶✓✔\uf0b7\uf0a7\uf076\uf0d8\uf0fc]\s*|[-*–—]\s+|\d{1,2}[.)]\s+)")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE_RE = re.compile(r"(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}\b")
URL_RE = re.compile(
    r"(?:https?://|www\.)\S+|\b(?:linkedin\.com|github\.com|behance\.net|dribbble\.com)/\S+"
    r"|\b[\w-]+\.(?:com|io|dev|me|org|net|co|app|page|site)(?:/\S*)?",
    re.I,
)
STATE_RE = re.compile(
    r"\b[A-Z][a-zA-Z.'-]+(?:\s[A-Z][a-zA-Z.'-]+){0,3},\s?"
    r"(?:A[LKZR]|C[AOT]|D[EC]|FL|GA|HI|I[DLNA]|K[SY]|LA|M[EDAINSOT]|N[EVHJMYCD]|O[HKR]|PA|RI|S[CD]|T[NX]|UT|V[TA]|W[AVIY])\b"
    r"|\bRemote\b"
)
MONTH = (r"(?:jan(?:uary)?|feb(?:ruary)?|mar(?:ch)?|apr(?:il)?|may|june?|july?|aug(?:ust)?|"
         r"sep(?:t(?:ember)?)?|oct(?:ober)?|nov(?:ember)?|dec(?:ember)?|spring|summer|fall|autumn|winter)\.?")
ONE_DATE = rf"(?:\b{MONTH},?\s+)?\b(?:19|20)\d{{2}}\b|\b\d{{1,2}}/(?:19|20)?\d{{2}}\b"
DATE_RE = re.compile(
    rf"(?:\b(?:expected|anticipated|graduat(?:ed|ing|ion)|class of)\s*:?\s*)?(?:{ONE_DATE})"
    rf"(?:\s*(?:-|–|—|to|until)\s*(?:{ONE_DATE}|present|current|now|ongoing|today))?"
    r"|\bgrades?\s*\d{1,2}\s*(?:-|–|—|to)\s*\d{1,2}\b"
    r"|\b\d{1,2}(?:th)?\s*(?:-|–|—|to)\s*\d{1,2}th\s+grade\b",
    re.I,
)
GPA_RE = re.compile(
    r"(?:\b(?:weighted|unweighted|cumulative|overall|major)\s+)?\bGPA\b\s*[:\-]?\s*(\d(?:\.\d{1,3})?)(\s*/\s*\d(?:\.\d+)?)?"
    r"(?:\s*\((?:un)?weighted\))?"
    r"|(\d\.\d{1,3})(\s*/\s*\d(?:\.\d+)?)?\s*(?:(?:un)?weighted\s+)?GPA\b",
    re.I,
)
TEST_RE = re.compile(r"\b(?:SAT|ACT|PSAT|PSAT/NMSQT)\b[^\n]*\d", re.I)
SCHOOL_RE = re.compile(r"\b(university|college|high school|academy|institute|school|polytechnic|conservatory)\b", re.I)
DEGREE_RE = re.compile(
    r"\b(bachelor\w*|master\w*|associate\w*|doctor\w*|ph\.?d|b\.s\.?|b\.a\.?|bs|ba|m\.s\.?|ms|mba|diploma|degree|"
    r"major|minor|ged|class of|candidate|concentration|ib)\b",
    re.I,
)
ROLE_RE = re.compile(
    r"\b(president|vice|captain|member|founder|co-founder|treasurer|secretary|officer|lead|leader|manager|intern|"
    r"associate|assistant|coordinator|director|editor|volunteer|tutor|instructor|coach|chair|representative|"
    r"delegate|analyst|engineer|developer|cashier|server|host|counselor|mentor|organizer|player|researcher|clerk)\b",
    re.I,
)
ORG_RE = re.compile(
    r"\b(inc|llc|ltd|corp|company|club|team|society|council|association|university|college|school|library|"
    r"hospital|foundation|restaurant|church|center|centre|group|lab|department|organization)\b\.?",
    re.I,
)
# Base forms that double as job titles (lead, tutor, volunteer, host...) are left out on purpose.
ACTION_VERBS = set("""
led manage managed build built create created organize organized organise ran won teach taught write wrote
helped worked assist assisted served train trained develop developed designed coordinate coordinated planned raised
mentored tutored launched increase increased reduce reduced improve improved analyze analyzed researched presented
handle handled greet greeted prepare prepared maintain maintained operate operated supported implement implemented
directed earned collaborate collaborated competed performed volunteered represented hosted
""".split())

SECTION_KEYS = [  # first match wins, so the specific ones come first
    ("tests", r"\b(test|tests|scores?|standardi[sz]ed|testing|sat|act)\b"),
    ("coursework", r"\b(coursework|course|courses|classes)\b"),
    ("volunteer", r"\b(volunteer\w*|community|service)\b"),
    ("activities", r"\b(extracurriculars?|extra|activities|activity|leadership|clubs?|organi[sz]ations?|involvement|athletics|sports)\b"),
    ("projects", r"\b(projects?|research|publications?|portfolio)\b"),
    ("awards", r"\b(awards?|honou?rs?|achievements?|accomplishments?|recognitions?|distinctions?)\b"),
    ("education", r"\b(education|academics?|academic|schooling)\b"),
    ("certifications", r"\b(certifications?|certificates?|licen[sc]es?)\b"),
    ("experience", r"\b(experiences?|employment|work|jobs?|career|history)\b"),
    ("skills", r"\b(skills?|technologies|tools|competencies|languages?|interests|proficiencies|strengths)\b"),
    ("summary", r"\b(summary|objective|profile|about|statement|overview)\b"),
]
HEADING_VOCAB = set("""
summary objective profile about me personal statement overview professional career education academic academics
background schooling experience experiences work employment history relevant jobs job extracurricular extracurriculars
extra curricular activities activity leadership clubs club organizations organization organisations involvement
athletics sports volunteer volunteering community service projects project technical awards award honors honours honor
achievements accomplishments recognition recognitions distinctions skills skill technologies tools competencies
languages language interests proficiencies strengths certifications certification licenses license licences
certificates coursework course courses classes test tests scores score standardized standardised testing sat act
other additional key selected research publications portfolio school
""".split())
ENTRY_SECTIONS = ("education", "experience", "activities", "volunteer", "projects")
LIST_SECTIONS = ("skills", "awards", "certifications", "coursework", "tests")


def _empty():
    result = {"profile": {"name": "", "email": "", "phone": "", "location": "", "links": [], "summary": ""}}
    for key in ENTRY_SECTIONS + LIST_SECTIONS:
        result[key] = []
    return result


def _clean_lines(text):
    lines = []
    for raw in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw.replace("\xa0", " ").replace("\u200b", "")
        line = re.sub(r"^#{1,6}\s*", "", line)          # markdown headings
        line = re.sub(r"\*\*|__", "", line)              # markdown bold
        line = re.sub(r"[ \f\v]{2,}", "  ", line).strip()
        lines.append(line)
    return lines


def _heading_key(line):
    raw = line.strip().strip(":|•-–— ").strip()
    if not raw or len(raw) > 48 or BULLET_RE.match(line):
        return None
    words = [w for w in re.split(r"[\s&/,\-+]+", raw.lower()) if w and w not in ("and", "of", "the", "my")]
    if not words or len(words) > 5 or any(w not in HEADING_VOCAB for w in words):
        return None
    joined = " ".join(words)
    for key, pattern in SECTION_KEYS:
        if re.search(pattern, joined):
            return key
    return None


def _looks_like_name(s):
    words = s.split()
    if not 2 <= len(words) <= 4 or not re.fullmatch(r"[A-Za-z .'\-]+", s):
        return False
    if any(w.lower().strip(".") in HEADING_VOCAB for w in words) or SCHOOL_RE.search(s) or ORG_RE.search(s):
        return False
    return all(w[0].isupper() for w in words)


def _cap(s):
    s = s.strip()
    return s[:1].upper() + s[1:] if s else s


def _tidy(s):
    return re.sub(r"\s{2,}", " ", s).strip(" ,;:|-–—()")


def _parse_header(lines, result, max_name_line=3):
    profile = result["profile"]
    leftovers = []
    seen = 0
    for line in lines:
        if not line:
            continue
        seen += 1
        rest = line
        for email in EMAIL_RE.findall(rest):
            profile["email"] = profile["email"] or email
        rest = EMAIL_RE.sub(" ", rest)
        for url in URL_RE.findall(rest):
            url = url.rstrip(".,;)")
            if url not in profile["links"]:
                profile["links"].append(url)
        rest = URL_RE.sub(" ", rest)
        for phone in PHONE_RE.findall(rest):
            profile["phone"] = profile["phone"] or phone.strip()
        rest = PHONE_RE.sub(" ", rest)
        loc = STATE_RE.search(rest)
        if loc and (rest != line or len(rest.split()) <= 5):
            profile["location"] = profile["location"] or loc.group(0)
            rest = rest[:loc.start()] + " " + rest[loc.end():]
        kept = []
        for piece in re.split(r"\s*(?:[|•·◆♦\t]|\s{2,})\s*", rest):
            piece = _tidy(piece)
            if not piece or not re.search(r"[A-Za-z]{2}", piece):
                continue
            if not profile["name"] and seen <= max_name_line and _looks_like_name(piece):
                profile["name"] = piece.title() if piece.isupper() else piece
            else:
                kept.append(piece)
        if kept:
            leftovers.append(" ".join(kept))
    return leftovers


def _pull_dates(text):
    best = None
    for m in DATE_RE.finditer(text):
        if best is None or len(m.group(0)) > len(best.group(0)):
            best = m
    if not best:
        return "", text
    dash = "–" if "grade" in best.group(0).lower() else " – "
    dates = re.sub(r"\s*(?:-|–|—|\bto\b|\buntil\b)\s*", dash, best.group(0).strip(), count=1)
    return dates, text[:best.start()] + " " + text[best.end():]


def _pull_location(text):
    m = STATE_RE.search(text)
    if not m:
        return "", text
    return m.group(0), text[:m.start()] + " " + text[m.end():]


def _segments(text):
    parts = re.split(r"\s*(?:\||•|·|\t|\s[–—-]\s|,\s+(?=[A-Z])|\s{2,}|\(|\))\s*", text)
    return [p for p in (_tidy(x) for x in parts) if p and re.search(r"[A-Za-z0-9]", p)]


def _split_role(seg):
    trail = r"\s+(?:since|from|in|during|starting|until|for)$"
    m = re.match(r"(.+?)\s+(?:at|@|for|with)\s+(.+)", seg)
    if m and len(m.group(1).split()) <= 5:
        return m.group(1), re.sub(trail, "", m.group(2))
    m = re.match(r"(.+?)\s+of\s+(.+)", seg)
    if m and len(m.group(1).split()) <= 3 and ROLE_RE.search(m.group(1)):
        return m.group(1), re.sub(trail, "", m.group(2))
    return seg, ""


def _is_detail(text, key):
    if GPA_RE.search(text) or (key == "education" and TEST_RE.search(text)):
        return True
    if re.match(r"^[A-Z][\w &/'-]{1,28}:\s+\S", text):
        return True
    words = re.findall(r"[A-Za-z0-9][\w'’.,+-]*", text)
    first = words[0].lower().strip(",") if words else ""
    if len(words) >= 4 and (first in ACTION_VERBS or (first.endswith("ed") and len(first) > 4)) and "|" not in text:
        return True
    headerish = ("|" in text or "\t" in text or DATE_RE.search(text)) and not text.rstrip().endswith(".")
    if headerish and len(words) <= 14:
        return False
    return len(words) >= 12 or (text.endswith(".") and len(words) >= 6)


def _split_entries(body, key):
    entries, cur, force_new = [], None, False
    for line in body:
        if not line:
            force_new = True
            continue
        is_bullet = bool(BULLET_RE.match(line))
        text = BULLET_RE.sub("", line).strip() if is_bullet else line.strip()
        if not text:
            continue
        if is_bullet or _is_detail(text, key):
            if cur is None:
                cur = {"header": [], "bullets": []}
                entries.append(cur)
            cur["bullets"].append(text)
        elif cur and cur["bullets"] and text[:1].islower() and not force_new:
            cur["bullets"][-1] += " " + text  # wrapped line
        else:
            header_has_date = cur is not None and any(DATE_RE.search(h) for h in cur["header"])
            new_school = (key == "education" and cur is not None and SCHOOL_RE.search(text) and not DEGREE_RE.search(text)
                          and any(SCHOOL_RE.search(h) for h in cur["header"]))
            if (cur is None or force_new or cur["bullets"] or len(cur["header"]) >= 3
                    or (header_has_date and len(cur["header"]) >= 2) or new_school):
                cur = {"header": [], "bullets": []}
                entries.append(cur)
            cur["header"].append(text)
        force_new = False
    return entries


def _entry(header_lines, bullets, key, result):
    text = " | ".join(header_lines)
    dates, text = _pull_dates(text)
    location, text = _pull_location(text)
    if key == "education":
        return _education(text, dates, location, bullets, result)
    segs = _segments(text)
    title, org = (segs[0], ", ".join(segs[1:])) if segs else ("", "")
    if title and not org:
        title, org = _split_role(title)
    if org and ((ORG_RE.search(title) and not ORG_RE.search(org)) or (ROLE_RE.search(org) and not ROLE_RE.search(title))):
        title, org = org, title
    return {"title": _cap(title), "org": org, "location": location, "dates": dates,
            "bullets": [_cap(b) for b in bullets if b.strip()]}


def _education(text, dates, location, bullets, result):
    gpa = ""
    lines_in = [text] + list(bullets)
    for i, line in enumerate(lines_in):
        m = GPA_RE.search(line)
        if m:
            value = m.group(1) or m.group(3)
            denom = (m.group(2) or m.group(4) or "").replace(" ", "")
            gpa = gpa or (value + denom)
            lines_in[i] = line[:m.start()] + " " + line[m.end():]
    text, bullets = lines_in[0], lines_in[1:]
    details = []
    for b in bullets:
        b = _tidy(b)
        if not b or not re.search(r"[A-Za-z]{3}", b):
            continue
        if STATE_RE.fullmatch(b):
            location = location or b
            continue
        cw =re.match(r"^(?:relevant\s+)?(?:course ?work|courses|classes)\s*[:\-]\s*(.+)", b, re.I)
        if cw:
            result["coursework"] += [c.strip(" .") for c in re.split(r"[,;]", cw.group(1)) if c.strip(" .")]
        elif TEST_RE.search(b):
            result["tests"].append(b)
        else:
            details.append(_cap(b))
    segs = _segments(text)
    school = next((s for s in segs if SCHOOL_RE.search(s) and not DEGREE_RE.search(s)), "")
    degree = next((s for s in segs if s != school and DEGREE_RE.search(s)), "")
    others = [s for s in segs if s not in (school, degree)]
    if not school and others:
        school = others.pop(0)
    if not degree and others:
        degree = others.pop(0)
    if others:
        degree = ", ".join([degree] + others) if degree else ", ".join(others)
    return {"school": school, "degree": degree, "location": location, "dates": dates, "gpa": gpa, "bullets": details}


def _clause_entry(text, key, result):
    """One free-typed line like 'President of Robotics Club, led 15 members (2023-2025)'."""
    dates, text = _pull_dates(text)
    location, text = _pull_location(text)
    text = re.sub(r"\(\s*\)", "", text).strip(" ,;")
    if key == "education":
        return _education(text, dates, location, [], result)
    words = text.split()
    if len(words) >= 8 and (words[0].lower() in ("i", "we") or words[0].lower() in ACTION_VERBS
                            or words[0].lower().endswith("ed")):
        return {"title": "", "org": "", "location": location, "dates": dates, "bullets": [_cap(text)]}
    parts = re.split(r"\s*(?:[,;:]|\s[–—-]\s|\|)\s*", text, maxsplit=1)
    title, org = _split_role(parts[0])
    bullets = [_cap(parts[1])] if len(parts) > 1 and parts[1].strip() else []
    return {"title": _cap(title), "org": org, "location": location, "dates": dates, "bullets": bullets}


def _list_items(body, key):
    items = []
    for line in body:
        if not line:
            continue
        is_bullet = bool(BULLET_RE.match(line))
        text = BULLET_RE.sub("", line).strip()
        if not text:
            continue
        if items and not is_bullet and text[:1].islower():
            items[-1] += " " + text
        elif key == "coursework":
            items += [c.strip(" .") for c in re.split(r"[,;]", re.sub(r"^(?:relevant\s+)?coursework\s*:\s*", "", text, flags=re.I)) if c.strip(" .")]
        else:
            items.append(text)
    return items


FREEFORM_RULES = [
    ("tests", r"\b(SAT|ACT|PSAT)\b.*\d"),
    ("education", r"\b(gpa|high school|university|college|graduat\w*|class of|major|diploma|degree|freshman|sophomore|junior|senior year)\b"),
    ("certifications", r"\b(certified|certification|certificate|cpr|first aid|license)\b"),
    ("awards", r"\b(award\w*|honor roll|honou?rs?|winner|won|\d(?:st|nd|rd|th) place|finalist|medal\w*|dean'?s list|scholar\b|recogni[sz]ed|semifinalist)\b"),
    ("volunteer", r"\b(volunteer\w*|community service|food bank|shelter|charity|nonprofit|non-profit|service hours)\b"),
    ("projects", r"\b(project|built an?|app\b|website|research(?:ed)?|prototype|hackathon)\b"),
    ("activities", r"\b(club|team|captain|president|vice president|treasurer|secretary|member|varsity|jv|band|orchestra|choir|"
                   r"student council|debate|robotics|model un|honor society|nhs|theat(?:er|re)|drama|sport\w*|football|soccer|"
                   r"basketball|baseball|softball|swim\w*|track|tennis|volleyball|lacrosse|cheer\w*|scouts?|yearbook|newspaper)\b"),
    ("experience", r"\b(work(?:ed|ing)?|job|intern\w*|employ\w*|cashier|server|barista|hired|part-time|full-time|company|"
                   r"inc\b|llc|manager|associate|lifeguard|babysit\w*|tutor\w*|retail|restaurant)\b"),
    ("skills", r"\b(skills?|proficient|fluent|bilingual|good at|familiar with|knowledge of|excel|photoshop|python|java|canva|typing)\b"),
]


def _freeform(lines, result):
    for raw in lines:
        if not raw:
            continue
        text = BULLET_RE.sub("", raw).strip()
        for piece in re.split(r"(?<=[.!?])\s+(?=[A-Z])|,?\s+and\s+(?=I\b)", text):
            piece = piece.strip().rstrip(".")
            # "I work as a barista at X" -> "barista at X", so the role and organization split cleanly
            piece = re.sub(r"^(?:hi|hello|hey)\b[,!]?\s*", "", piece, flags=re.I)
            piece = re.sub(r"^(?:i|we)\s+(?:work|worked|am|was|serve|served|have been)\s+(?:as\s+)?(?:an?\s+|the\s+)?",
                           "", piece, flags=re.I)
            piece = re.sub(r"\b(of|at|for)\s+the\s+", r"\1 ", piece)
            if not piece:
                continue
            key = next((k for k, pat in FREEFORM_RULES if re.search(pat, piece, re.I)), None)
            if key is None:
                profile = result["profile"]
                profile["summary"] = (profile["summary"] + " " + piece + ".").strip()
            elif key in LIST_SECTIONS:
                item = re.sub(r"^(?:my\s+)?(?:skills|awards|certifications|courses)\s*(?:are|include|:)\s*", "", piece, flags=re.I)
                result[key].append(_cap(item))
            elif key == "education":
                gpa = GPA_RE.search(piece)
                school = re.search(r"((?:[A-Z][\w.'&-]*\s+){0,5}(?:University|College|High School|Academy|Institute|School)"
                                   r"(?:\s+of\s+[A-Z][\w.'-]*(?:\s+[A-Z][\w.'-]*)*)?)", piece)
                if school or not result["education"]:
                    dates, _ = _pull_dates(piece)
                    degree = DEGREE_RE.search(piece)
                    result["education"].append({
                        "school": school.group(1).strip() if school else "", "degree": "", "location": "",
                        "dates": dates, "gpa": "", "bullets": [] if (school or gpa) else [_cap(piece)],
                    })
                    if degree and not school:
                        result["education"][-1]["degree"] = _cap(piece)
                if gpa:
                    result["education"][-1]["gpa"] = (gpa.group(1) or gpa.group(3)) + (gpa.group(2) or gpa.group(4) or "").replace(" ", "")
            else:
                result[key].append(_clause_entry(piece, key, result))


def parse_resume_text(text):
    lines = _clean_lines(text)
    blocks, current, found = {"header": []}, "header", False
    for line in lines:
        key = _heading_key(line) if line else None
        if key:
            found, current = True, key
            blocks.setdefault(key, [])
            continue
        blocks.setdefault(current, []).append(line)

    result = _empty()
    header = blocks.pop("header")
    if not found:
        intro = re.search(r"\b(?i:my name is|i am|i'm)\s+([A-Z][a-z'-]+(?:\s+[A-Z][a-z'-]+){1,2})\b", text)
        if intro:
            result["profile"]["name"] = intro.group(1)
        leftovers = _parse_header(header[:4], result, max_name_line=2)
        rest = header[4:]
        # Lines in the top block that weren't contact info are real content in free-typed text.
        _freeform(leftovers + rest, result)
    else:
        leftovers = _parse_header(header, result)
        long_lines = [l for l in leftovers if len(l.split()) >= 7]
        if long_lines:
            result["profile"]["summary"] = " ".join(long_lines)
        for key, body in blocks.items():
            if key == "summary":
                text = " ".join(BULLET_RE.sub("", l).strip() for l in body if l.strip())
                result["profile"]["summary"] = (result["profile"]["summary"] + " " + text).strip()
            elif key in ENTRY_SECTIONS:
                for raw in _split_entries(body, key):
                    if not raw["header"] and key != "education" and len(raw["bullets"]) > 0:
                        result[key] += [_clause_entry(b, key, result) for b in raw["bullets"]]
                    else:
                        result[key].append(_entry(raw["header"], raw["bullets"], key, result))
            else:
                result[key] += _list_items(body, key)

    result["profile"]["links"] = " | ".join(result["profile"]["links"])
    return result
