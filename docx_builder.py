"""Write the resume model the browser sends into a Word document."""

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONTS = {"sans": "Calibri", "serif": "Georgia", "modern": "Arial"}
INK, MUTED = "1F1F1F", "5F5F5F"
MARGIN = 0.55
# Elements that must follow <w:pBdr> inside <w:pPr>, or Word calls the file corrupt.
AFTER_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
              "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
              "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
              "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
              "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _hex(color, fallback="1A1A1A"):
    c = (color or "").lstrip("#")
    return c.upper() if re.fullmatch(r"[0-9a-fA-F]{6}", c) else fallback


def _clamp(value, lo, hi, default):
    try:
        return max(lo, min(hi, float(value)))
    except (TypeError, ValueError):
        return default


def _border(paragraph, color, size=8, side="bottom"):
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    edge = OxmlElement(f"w:{side}")
    for key, val in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "2"), ("w:color", color)):
        edge.set(qn(key), val)
    bdr.append(edge)
    ppr.insert_element_before(bdr, *AFTER_PBDR)


def _run(p, text, size=None, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _labelled(p, text, size=None):
    """'Languages: Python, Java' -> bold label, plain rest."""
    m = re.match(r"^([^:]{1,32}):\s*(.+)$", text)
    if m:
        _run(p, m.group(1) + ": ", size, bold=True)
        _run(p, m.group(2), size)
    else:
        _run(p, text, size)


def _bullet(doc, text, gap):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.left_indent, pf.first_line_indent = Inches(0.22), Inches(-0.16)
    pf.space_before, pf.space_after = Pt(0.8 * gap), Pt(0)
    _labelled(p, text)
    return p


def _entry(doc, item, size, gap, width):
    heading, sub = item.get("heading", ""), item.get("sub", "")
    dates, meta, loc = item.get("dates", ""), item.get("meta", ""), item.get("location", "")
    if heading or sub or dates:
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        pf.space_before, pf.keep_with_next = Pt(5 * gap), True
        if heading:
            _run(p, heading, bold=True)
        if sub:
            _run(p, (", " if heading else "") + sub, italic=bool(heading), bold=not heading)
        if dates:
            _run(p, "\t" + dates, size * 0.95, color=MUTED)
    if meta or loc:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.keep_with_next = True
        if meta:
            _run(p, meta, size * 0.95, color=MUTED)
        if loc:
            _run(p, "\t" + loc, size * 0.95, italic=True, color=MUTED)
    for b in item.get("bullets", []):
        _bullet(doc, b, gap)


def build_docx(model, design):
    template = design.get("template", "clean")
    accent = _hex(design.get("accent"))
    head_color = INK if template == "classic" else (MUTED if template == "minimal" else accent)
    font = FONTS.get(design.get("font"), "Calibri")
    size = _clamp(design.get("size"), 8, 13, 10.5)
    gap = _clamp(design.get("gap"), 0.4, 1.5, 1.0)
    center = template == "classic"

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Inches(MARGIN)
    width = Inches(8.5 - 2 * MARGIN)

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(size)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.08

    p = doc.add_paragraph()
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_color = accent if template in ("clean", "bold") else INK
    _run(p, model.get("name") or "Your Name", size * 2.3, bold=True, color=name_color)

    contact = [c for c in model.get("contact", []) if c]
    if contact:
        p = doc.add_paragraph()
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p, "  •  ".join(contact), size * 0.95, color=MUTED)
    p.paragraph_format.space_after = Pt(4)
    if template == "bold":
        _border(p, accent, 24)
    elif template == "classic":
        _border(p, INK, 6)

    for sec in model.get("sections", []):
        h = doc.add_paragraph()
        hf = h.paragraph_format
        hf.space_before, hf.space_after, hf.keep_with_next = Pt(size * 1.05 * gap), Pt(3 * gap), True
        title = sec.get("title", "")
        _run(h, title if template == "bold" else title.upper(), size * (1.2 if template == "bold" else 1.08),
             bold=True, color=head_color)
        if template in ("clean", "classic"):
            _border(h, head_color, 6)

        kind, items = sec.get("type"), sec.get("items", [])
        if kind == "text":
            doc.add_paragraph(sec.get("text", ""))
        elif kind == "entries":
            for item in items:
                _entry(doc, item, size, gap, width)
        elif sec.get("style") == "inline":
            doc.add_paragraph("  ·  ".join(items))
        elif sec.get("style") == "skills":
            for line in items:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0.8 * gap)
                _labelled(p, line)
        else:
            for line in items:
                _bullet(doc, line, gap)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
